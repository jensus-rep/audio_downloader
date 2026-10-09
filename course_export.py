"""Erfasst die Inhalte eines Moodle-Kurses: Inventar, Texte, Dateien und Medien.

Ziel ist ein vollständiges Bild des Kurses ohne blinde Flecken: Jede Aktivität
wird im Inventar geführt, auch wenn ihr Inhalt nicht geholt werden kann oder
bewusst nicht geöffnet wird (z. B. Tests, die nur einmal aufrufbar sind).
"""

from __future__ import annotations

import base64
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable
from urllib.parse import parse_qs, urldefrag, urljoin, urlparse

from playwright.sync_api import BrowserContext, Page

from config import Config
from crawler import AudioCrawler, LoginError
from downloader import AudioDownloader


LogFn = Callable[[str], None]

# Activities whose content is read by opening their page.
TEXT_TYPES = {"page", "book", "h5pactivity", "hvp", "glossary", "wiki", "imscp"}
# Activities whose content is a file or a list of files.
FILE_TYPES = {"resource", "folder"}
# Activities that are only listed. Opening them could start an attempt, create
# a submission or otherwise change state; the final test is a one-shot quiz.
NEVER_OPEN = {
    "quiz": "Test: wird nie geöffnet, damit kein Versuch beginnt",
    "assign": "Abgabe: wird nicht geöffnet",
    "forum": "Forum: wird nicht geöffnet",
    "feedback": "Umfrage: wird nicht geöffnet",
    "questionnaire": "Umfrage: wird nicht geöffnet",
    "survey": "Umfrage: wird nicht geöffnet",
    "choice": "Abstimmung: wird nicht geöffnet",
    "lesson": "Lektion: wird nicht geöffnet, weil sie Antworten speichert",
    "scorm": "Lernpaket: wird nicht geöffnet, weil es Versuche zählt",
    "lti": "Externes Werkzeug: wird nicht geöffnet",
    "workshop": "Peer-Review: wird nicht geöffnet",
    "chat": "Chat: wird nicht geöffnet",
    "data": "Datenbank: wird nicht geöffnet",
    "attendance": "Anwesenheit: wird nicht geöffnet",
    "bigbluebuttonbn": "Videokonferenz: wird nicht geöffnet",
    "zoom": "Videokonferenz: wird nicht geöffnet",
}
FORBIDDEN_URL_PARTS = (
    "logout",
    "sesskey=",
    "/mod/quiz/",
    "/mod/assign/",
    "/mod/lesson/",
    "/mod/scorm/",
    "/mod/forum/post",
    "/course/edit",
    "/enrol/",
)
TEXT_FILE_SUFFIXES = {".pdf", ".pptx", ".docx", ".txt", ".md"}
MEDIA_SUFFIXES = {".mp3", ".wav", ".m4a", ".ogg", ".mp4", ".webm", ".mov", ".m4v"}
LV_PATTERN = re.compile(
    r"\b(?:LV|Lehrveranstaltung|Lehreinheit|Lerneinheit)\s*0*(\d{1,3})\b",
    re.IGNORECASE,
)

COURSE_STRUCTURE_JS = r"""
() => {
  const text = el => (el ? el.innerText.replace(/\s+\n/g, "\n").trim() : "");
  const nameOf = li => {
    if (li.dataset.activityname) return li.dataset.activityname.trim();
    const holder = li.querySelector(".instancename, .activityname");
    if (!holder) return "";
    const clone = holder.cloneNode(true);
    clone.querySelectorAll(".accesshide, .sr-only, .visually-hidden").forEach(n => n.remove());
    return clone.innerText.trim();
  };
  const mediaIn = root => {
    const found = [];
    root.querySelectorAll("iframe[src], video[src], audio[src], source[src], embed[src], object[data]")
      .forEach(el => found.push({ tag: el.tagName.toLowerCase(), src: el.src || el.data || el.getAttribute("src") || el.getAttribute("data") }));
    root.querySelectorAll('a[href*="pluginfile.php"]').forEach(a => found.push({ tag: "a", src: a.href, text: a.innerText.trim() }));
    return found;
  };
  const sectionOf = el => el.closest('li.section, li.course-section, [data-for="section"]');
  const sectionInfo = sec => {
    if (!sec) return { number: null, name: "", summary: "" };
    const nameEl = sec.querySelector('.sectionname, [data-for="section_title"], h3');
    const summaryEl = sec.querySelector(':scope .summarytext, :scope .summary, :scope [data-for="sectioninfo"] .summarytext');
    let number = sec.dataset.number ?? sec.dataset.sectionnumber ?? null;
    if (number === null && sec.id && /^section-\d+$/.test(sec.id)) number = sec.id.split("-")[1];
    return { number: number === null ? null : Number(number), name: text(nameEl), summary: text(summaryEl) };
  };
  const sections = [...document.querySelectorAll('li.section, li.course-section, [data-for="section"]')]
    .filter(sec => !sec.parentElement.closest('li.section, li.course-section, [data-for="section"]'))
    .map(sectionInfo);
  const activities = [...document.querySelectorAll("li.activity")].map(li => {
    const typeClass = [...li.classList].find(c => c.startsWith("modtype_"));
    const link = li.querySelector("a.aalink, .activityname a, a.activity-link, a[href*='/mod/']");
    const idMatch = (li.id || "").match(/module-(\d+)/);
    const inline = li.querySelector(".contentwithoutlink, .activity-altcontent, .activity-description, .description, .no-overflow");
    const availability = li.querySelector(".availabilityinfo, .activity-information [data-region='availabilityinfo']");
    return {
      cmid: idMatch ? idMatch[1] : (li.dataset.id || null),
      modtype: typeClass ? typeClass.slice("modtype_".length) : "",
      name: nameOf(li),
      url: link ? link.href : "",
      inline_text: text(inline),
      availability: text(availability),
      hidden: li.classList.contains("dimmed") || !!li.querySelector(".dimmed, .dimmed_text"),
      media: mediaIn(li),
      section: sectionInfo(sectionOf(li)),
    };
  });
  return { title: document.title, heading: text(document.querySelector(".page-header-headings h1, header h1, h1")), sections, activities };
}
"""

# Converts the main content of a page to light Markdown: headings, paragraphs,
# list items and table rows keep their structure; everything else is plain text.
PAGE_TO_MARKDOWN_JS = r"""
(rootSelectors) => {
  let root = null;
  for (const sel of rootSelectors) { root = document.querySelector(sel); if (root) break; }
  root = root || document.body;
  const skip = new Set(["SCRIPT", "STYLE", "NOSCRIPT", "TEMPLATE", "NAV", "BUTTON", "SELECT"]);
  const lines = [];
  const clean = s => s.replace(/[ \t ]+/g, " ").trim();
  const walk = node => {
    if (node.nodeType !== 1 || skip.has(node.tagName)) return;
    const el = node;
    if (el.matches(".accesshide, .sr-only, .visually-hidden, .skip-block, .activity-navigation, .modified")) return;
    const tag = el.tagName;
    if (/^H[1-6]$/.test(tag)) { const t = clean(el.innerText); if (t) lines.push("", "#".repeat(Number(tag[1])) + " " + t, ""); return; }
    if (tag === "LI") { const t = clean(el.innerText).replace(/\n+/g, " "); if (t) lines.push("- " + t); return; }
    if (tag === "TR") { const cells = [...el.children].map(c => clean(c.innerText).replace(/\n+/g, " ")); if (cells.some(Boolean)) lines.push("| " + cells.join(" | ") + " |"); return; }
    if (tag === "P" || tag === "BLOCKQUOTE" || tag === "PRE" || tag === "FIGCAPTION" || tag === "DT" || tag === "DD") {
      const t = el.innerText.split("\n").map(clean).filter(Boolean).join("\n"); if (t) lines.push("", (tag === "BLOCKQUOTE" ? "> " : "") + t, ""); return;
    }
    const blockChildren = [...el.children].some(c => /^(DIV|P|H[1-6]|UL|OL|LI|TABLE|TBODY|THEAD|TR|SECTION|ARTICLE|BLOCKQUOTE|PRE|FIGURE|DL)$/.test(c.tagName));
    if (!blockChildren) { const t = el.innerText.split("\n").map(clean).filter(Boolean).join("\n"); if (t) lines.push("", t, ""); return; }
    [...el.children].forEach(walk);
  };
  walk(root);
  const markdown = lines.join("\n").replace(/\n{3,}/g, "\n\n").trim();
  const links = [...root.querySelectorAll("a[href]")].map(a => ({ text: clean(a.innerText), href: a.href })).filter(l => l.href && !l.href.startsWith("javascript:"));
  const media = [];
  root.querySelectorAll("iframe[src], video[src], audio[src], source[src], embed[src], object[data]")
    .forEach(el => media.push({ tag: el.tagName.toLowerCase(), src: el.src || el.data || el.getAttribute("src") || el.getAttribute("data") }));
  return { markdown, links, media, title: document.title };
}
"""
MAIN_SELECTORS = ['[role="main"]', "#region-main", "main", "#page-content"]


@dataclass
class Activity:
    index: int
    cmid: str | None
    modtype: str
    name: str
    url: str
    section_number: int | None
    section_name: str
    lv_number: int | None = None
    availability: str = ""
    hidden: bool = False
    status: str = "gelistet"
    note: str = ""
    text_file: str | None = None
    files: list[str] = field(default_factory=list)
    file_texts: list[str] = field(default_factory=list)
    external_links: list[str] = field(default_factory=list)
    media: list[dict] = field(default_factory=list)
    characters: int = 0


@dataclass
class ExportSummary:
    activities: int
    with_text: int
    files: int
    listed_only: int
    failed: int
    lv_numbers: list[int]
    export_dir: Path

    def print(self, log: LogFn = print) -> None:
        log("")
        log("Zusammenfassung")
        log(f"Aktivitäten im Kurs: {self.activities}")
        log(f"Mit erfasstem Text: {self.with_text}")
        log(f"Heruntergeladene Dateien: {self.files}")
        log(f"Nur gelistet (bewusst nicht geöffnet oder ohne Inhalt): {self.listed_only}")
        log(f"Fehler: {self.failed}")
        if self.lv_numbers:
            log("LV-Nummern im Kurs: " + ", ".join(str(n) for n in self.lv_numbers))
        log(f"Export-Ordner: {self.export_dir}")


def lv_number(*texts: str) -> int | None:
    for value in texts:
        match = LV_PATTERN.search(value or "")
        if match:
            return int(match.group(1))
    return None


def slug(value: str, limit: int = 60) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_only = normalized.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", ascii_only).strip("-").lower()
    return (cleaned or "ohne-titel")[:limit].rstrip("-")


def is_safe_to_open(url: str, origin: str) -> bool:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or parsed.netloc != origin:
        return False
    lowered = url.lower()
    return not any(part in lowered for part in FORBIDDEN_URL_PARTS)


class CourseExporter:
    def __init__(
        self,
        context: BrowserContext,
        config: Config,
        log: LogFn = print,
        with_media: bool = False,
    ) -> None:
        self.context = context
        self.config = config
        self.log = log
        self.with_media = with_media
        self.crawler = AudioCrawler(context=context, config=config, log=log)
        self.page: Page = self.crawler.page
        # The audio crawler inspects every network response; the export does not
        # need that, and on media-heavy courses it exhausted memory.
        self.page.remove_listener("response", self.crawler._capture_audio_response)
        self.vimeo: dict[str, dict] = {}

    # -- entry point -------------------------------------------------------

    def export_course(self, course_url: str) -> ExportSummary:
        self.log("Öffne Kursübersicht")
        self.crawler.open_authenticated(course_url)
        origin = urlparse(self.page.url).netloc
        course_id = parse_qs(urlparse(course_url).query).get("id", [None])[0]

        structure = self._read_structure(self.page)
        title = structure["heading"] or structure["title"] or f"Kurs {course_id}"
        root = self.config.export_dir / f"kurs-{course_id or slug(title)}"
        (root / "texte").mkdir(parents=True, exist_ok=True)
        self.root = root

        section_urls = self.crawler._extract_course_section_urls(course_id)
        self.log(f"{len(section_urls)} Kursabschnitte gefunden")
        sections: dict[int | None, dict] = {}
        raw_activities: dict[str, dict] = {}
        pages = [urldefrag(course_url).url, *section_urls]
        for index, page_url in enumerate(dict.fromkeys(pages), start=1):
            if index > 1:
                self.log(f"Lese Kursabschnitt {index - 1} von {len(section_urls)}")
                self._open(page_url)
                structure = self._read_structure(self.page)
            for section in structure["sections"]:
                known = sections.get(section["number"])
                if not known or len(section["summary"]) > len(known["summary"]):
                    sections[section["number"]] = section
            for item in structure["activities"]:
                key = item["cmid"] or item["url"] or f"{item['modtype']}:{item['name']}"
                known = raw_activities.get(key)
                if not known or len(item["inline_text"]) > len(known["inline_text"]):
                    raw_activities[key] = item
            if self.with_media:
                self._capture_vimeo(structure["activities"])

        if self.vimeo:
            (root / "medien" / "vimeo.json").write_text(
                json.dumps(self.vimeo, ensure_ascii=False, indent=1), encoding="utf-8"
            )
            self.log(f"Vimeo: {sum(1 for v in self.vimeo.values() if v['art'] == 'untertitel')} Untertitel, "
                     f"{sum(1 for v in self.vimeo.values() if v['art'] == 'tonspur')} Tonspuren, "
                     f"{sum(1 for v in self.vimeo.values() if v['art'] == 'fehler')} Fehler")

        activities = [
            self._activity_from(position, item)
            for position, item in enumerate(raw_activities.values(), start=1)
        ]
        name_activities(activities, sections, list(raw_activities.values()))
        self.log(f"{len(activities)} Aktivitäten gefunden")
        self._write_section_summaries(sections)

        for activity, item in zip(activities, raw_activities.values()):
            try:
                self._capture(activity, item, origin)
            except LoginError:
                raise
            except Exception as exc:
                activity.status = "fehler"
                activity.note = str(exc)[:300]
                self.log(f"Fehler bei „{activity.name}“: {activity.note}")

        self._write_inventory(title, course_url, sections, activities)
        lv_numbers = sorted({a.lv_number for a in activities if a.lv_number is not None})
        return ExportSummary(
            activities=len(activities),
            with_text=sum(1 for a in activities if a.text_file or a.file_texts),
            files=sum(len(a.files) for a in activities),
            listed_only=sum(1 for a in activities if a.status in {"gelistet", "nicht geöffnet", "gesperrt"}),
            failed=sum(1 for a in activities if a.status == "fehler"),
            lv_numbers=lv_numbers,
            export_dir=root,
        )

    # -- structure ---------------------------------------------------------

    def _read_structure(self, page: Page) -> dict:
        return page.evaluate(COURSE_STRUCTURE_JS)

    def _activity_from(self, position: int, item: dict) -> Activity:
        section = item.get("section") or {}
        activity = Activity(
            index=position,
            cmid=item.get("cmid"),
            modtype=item.get("modtype") or "unbekannt",
            name=item.get("name") or "(ohne Titel)",
            url=item.get("url") or "",
            section_number=section.get("number"),
            section_name=section.get("name") or "",
            availability=item.get("availability") or "",
            hidden=bool(item.get("hidden")),
            media=list(item.get("media") or []),
        )
        activity.lv_number = lv_number(activity.name, activity.section_name)
        return activity

    # -- capture per activity ---------------------------------------------

    def _capture(self, activity: Activity, item: dict, origin: str) -> None:
        inline = (item.get("inline_text") or "").strip()
        if activity.modtype in NEVER_OPEN:
            activity.status = "nicht geöffnet"
            activity.note = NEVER_OPEN[activity.modtype]
            if inline:
                self._save_text(activity, inline, source="Beschreibung auf der Kursseite")
            return
        if inline and activity.modtype in {"label", "subsection"} | TEXT_TYPES | FILE_TYPES | {"url"}:
            # Labels carry their whole content on the course page; for other
            # types the course page may show a description worth keeping.
            self._save_text(activity, inline, source="Text auf der Kursseite")
            activity.status = "text"
        if activity.modtype in {"label", "subsection"}:
            if not inline:
                activity.status = "ohne Inhalt"
            return
        if not activity.url:
            activity.status = "gesperrt" if activity.availability else "gelistet"
            activity.note = activity.availability or "Kein Link auf der Kursseite"
            return
        if not is_safe_to_open(activity.url, origin):
            activity.status = "nicht geöffnet"
            activity.note = "Link führt aus dem Kurs heraus oder ist ausgeschlossen"
            activity.external_links.append(activity.url)
            return

        self.log(f"Erfasse {activity.modtype}: {activity.name}")
        if activity.modtype == "resource":
            self._capture_resource(activity, origin)
        elif activity.modtype == "folder":
            self._capture_folder(activity, origin)
        elif activity.modtype == "url":
            self._capture_url(activity)
        elif activity.modtype == "book":
            self._capture_book(activity, origin)
        elif activity.modtype in TEXT_TYPES:
            self._capture_page(activity, activity.url)
        else:
            activity.status = "nicht geöffnet"
            activity.note = f"Unbekannter Aktivitätstyp „{activity.modtype}“, sicherheitshalber nicht geöffnet"
        if self.with_media:
            self._download_media(activity, origin)

    def _capture_page(self, activity: Activity, url: str, source: str = "Seite") -> None:
        self._open(url)
        result = self.page.evaluate(PAGE_TO_MARKDOWN_JS, MAIN_SELECTORS)
        markdown = result["markdown"]
        frame_texts = self._frame_texts()
        if frame_texts:
            markdown += "\n\n## Eingebettete Inhalte\n\n" + "\n\n".join(frame_texts)
        self._save_text(activity, markdown, source=source)
        activity.media.extend(result["media"])
        activity.status = "text" if markdown.strip() else "ohne Text"
        for link in result["links"]:
            href = link["href"]
            if "pluginfile.php" in href:
                self._download_file(activity, href)
            elif urlparse(href).netloc and urlparse(href).netloc != urlparse(self.page.url).netloc:
                activity.external_links.append(href)

    def _capture_book(self, activity: Activity, origin: str) -> None:
        cmid = activity.cmid or parse_qs(urlparse(activity.url).query).get("id", [None])[0]
        scheme = urlparse(activity.url).scheme or "https"
        print_url = f"{scheme}://{origin}/mod/book/tool/print/index.php?id={cmid}"
        try:
            self._capture_page(activity, print_url, source="Buch (Druckansicht, alle Kapitel)")
            if activity.characters > 200:
                return
        except Exception:
            pass
        self._capture_page(activity, activity.url, source="Buch")
        chapters = [
            href
            for href in self.page.eval_on_selector_all(
                'a[href*="chapterid="]', "els => [...new Set(els.map(a => a.href))]"
            )
            if is_safe_to_open(href, origin)
        ]
        for href in chapters:
            self._capture_page(activity, href, source="Buchkapitel")

    def _capture_resource(self, activity: Activity, origin: str) -> None:
        url = activity.url + ("&" if "?" in activity.url else "?") + "redirect=1"
        if self._download_file(activity, url):
            activity.status = "datei" if activity.files else "medium"
            return
        # Resource shown embedded in a page: collect the file links from there.
        self._open(activity.url)
        links = self.page.eval_on_selector_all(
            'a[href*="pluginfile.php"], object[data*="pluginfile.php"], iframe[src*="pluginfile.php"], embed[src*="pluginfile.php"]',
            "els => [...new Set(els.map(e => e.href || e.data || e.src))]",
        )
        for href in links:
            self._download_file(activity, href)
        activity.status = "datei" if activity.files else "medium" if activity.media else "ohne Datei"

    def _capture_folder(self, activity: Activity, origin: str) -> None:
        self._open(activity.url)
        links = self.page.eval_on_selector_all(
            'a[href*="pluginfile.php"]', "els => [...new Set(els.map(a => a.href))]"
        )
        for href in links:
            self._download_file(activity, href)
        activity.status = "datei" if activity.files else "medium" if activity.media else "ohne Datei"

    def _capture_url(self, activity: Activity) -> None:
        url = activity.url + ("&" if "?" in activity.url else "?") + "redirect=1"
        response = self.context.request.get(
            url, max_redirects=0, fail_on_status_code=False, timeout=self.config.page_timeout_ms
        )
        target = response.headers.get("location")
        if not target:
            self._open(activity.url)
            target = self.page.eval_on_selector_all(
                ".urlworkaround a, [role='main'] a[href^='http']",
                "els => els.length ? els[0].href : ''",
            )
        if target:
            activity.external_links.append(urljoin(activity.url, target))
            activity.status = "link"
        else:
            activity.status = "ohne Link"

    # -- files and media ---------------------------------------------------

    def _download_file(self, activity: Activity, url: str) -> bool:
        response = self.context.request.get(
            url, timeout=self.config.download_timeout_ms, fail_on_status_code=False
        )
        content_type = response.headers.get("content-type", "").split(";")[0].strip().lower()
        disposition = response.headers.get("content-disposition", "")
        if not response.ok or content_type in {"text/html", "application/xhtml+xml"}:
            return False
        name = AudioDownloader._filename_from_content_disposition(disposition) or Path(
            urlparse(response.url).path
        ).name or "datei"
        name = AudioDownloader._sanitize_filename(name)
        suffix = Path(name).suffix.lower()
        if suffix in MEDIA_SUFFIXES and not self.with_media:
            activity.media.append({"tag": "datei", "src": url, "text": name})
            return True
        folder = self.root / ("medien" if suffix in MEDIA_SUFFIXES else "dateien")
        folder.mkdir(parents=True, exist_ok=True)
        destination = folder / f"{activity.cmid or activity.index}_{name}"
        relative = destination.relative_to(self.root).as_posix()
        if relative in activity.files:
            return True
        if not destination.exists():
            AudioDownloader._write_atomic(destination, response.body())
        activity.files.append(relative)
        if suffix in TEXT_FILE_SUFFIXES:
            text = extract_file_text(destination)
            if text:
                text_path = self.root / "texte" / "dateien" / f"{destination.stem}.md"
                text_path.parent.mkdir(parents=True, exist_ok=True)
                text_path.write_text(
                    self._front_matter(activity, f"Datei {name}") + text + "\n", encoding="utf-8"
                )
                activity.file_texts.append(text_path.relative_to(self.root).as_posix())
                activity.characters += len(text)
        return True

    def _download_media(self, activity: Activity, origin: str) -> None:
        for item in list(activity.media):
            src = item.get("src") or ""
            if urlparse(src).netloc == origin and Path(urlparse(src).path).suffix.lower() in MEDIA_SUFFIXES:
                self._download_file(activity, src)

    def _capture_vimeo(self, items: list[dict]) -> None:
        """Save captions (preferred) or the audio track of embedded Vimeo videos on the current page.

        Works only inside the logged-in course page: embed-restricted videos refuse requests
        from anywhere else. Captions come from the player's text tracks; without captions the
        DASH audio segments are joined into one .m4a file for transcription.
        """
        titles = {}
        for item in items:
            name = item.get("name") or ""
            if not name or name == "(ohne Titel)":
                name = (item.get("inline_text") or "").strip().split("\n")[0][:80]
            for media in item.get("media") or []:
                match = re.search(r"player\.vimeo\.com/video/(\d+)", media.get("src") or "")
                if match:
                    titles.setdefault(match.group(1), name)
        folder = self.root / "medien"
        for frame in self.page.frames:
            match = re.search(r"player\.vimeo\.com/video/(\d+)", frame.url)
            if not match or match.group(1) in self.vimeo:
                continue
            vid = match.group(1)
            title = titles.get(vid) or vid
            entry = {"titel": title, "art": "fehler", "datei": None}
            self.vimeo[vid] = entry
            try:
                frame.wait_for_load_state("domcontentloaded", timeout=15_000)
                raw = frame.evaluate("() => window.playerConfig ? JSON.stringify(window.playerConfig) : null")
                if not raw:
                    raw = frame.evaluate(
                        "async id => { const r = await fetch(`/video/${id}/config`); return r.ok ? await r.text() : null; }",
                        vid,
                    )
                cfg = json.loads(raw or "{}")
                request = cfg.get("request", {})
                entry["dauer_s"] = cfg.get("video", {}).get("duration")
                folder.mkdir(parents=True, exist_ok=True)
                stem = f"vimeo_{vid}_{slug(title, 50)}"
                tracks = request.get("text_tracks") or []
                # Only German captions: auto captions in another language for German speech are unusable,
                # the audio track is the better source then.
                german = [t for t in tracks if (t.get("lang") or "").startswith("de")]
                headers = {"Referer": "https://player.vimeo.com/"}
                if german:
                    url = urljoin("https://player.vimeo.com/", german[0]["url"])
                    body = self.context.request.get(url, headers=headers, timeout=self.config.download_timeout_ms).body()
                    (folder / f"{stem}.vtt").write_bytes(body)
                    entry.update(art="untertitel", datei=f"medien/{stem}.vtt", sprache=german[0].get("lang"))
                else:
                    dash = request.get("files", {}).get("dash") or {}
                    cdn = (dash.get("cdns") or {}).get(dash.get("default_cdn")) or next(iter((dash.get("cdns") or {}).values()), {})
                    playlist_url = cdn.get("url") or cdn.get("avc_url")
                    if not playlist_url:
                        raise RuntimeError("weder Untertitel noch Audiospur angeboten")
                    playlist = json.loads(self.context.request.get(playlist_url, headers=headers).body())
                    audio = sorted(playlist.get("audio") or [], key=lambda a: a.get("bitrate", 0))
                    if not audio:
                        raise RuntimeError("keine Audiospur in der Playlist")
                    track = audio[-1]
                    base = urljoin(urljoin(playlist_url, playlist.get("base_url", "")), track.get("base_url", ""))
                    data = bytearray(base64.b64decode(track["init_segment"]))
                    for segment in track["segments"]:
                        data += self.context.request.get(urljoin(base, segment["url"]), headers=headers).body()
                    AudioDownloader._write_atomic(folder / f"{stem}.m4a", bytes(data))
                    entry.update(art="tonspur", datei=f"medien/{stem}.m4a")
                self.log(f"Vimeo {entry['art']}: {title}")
            except Exception as exc:
                entry["fehler"] = str(exc)[:300]
                self.log(f"Vimeo-Fehler bei „{title}“: {entry['fehler']}")

    def _frame_texts(self) -> list[str]:
        texts = []
        main_origin = urlparse(self.page.url).netloc
        for frame in self.page.frames[1:]:
            if urlparse(frame.url).netloc != main_origin:
                continue
            try:
                value = frame.locator("body").inner_text(timeout=3_000).strip()
            except Exception:
                continue
            if value:
                texts.append(value)
        return texts

    # -- output ------------------------------------------------------------

    def _open(self, url: str) -> None:
        self.page.wait_for_timeout(self.config.request_delay_ms)
        self.crawler._goto(self.page, url)
        if self.crawler._login_form_visible(self.page):
            raise LoginError("Die Sitzung ist beim Erfassen des Kurses abgelaufen.")
        self.crawler._wait_for_dynamic_content()

    def _front_matter(self, activity: Activity, source: str) -> str:
        lines = [
            f"# {activity.name}",
            "",
            f"- Typ: {activity.modtype}",
            f"- Abschnitt: {activity.section_number if activity.section_number is not None else '–'} {activity.section_name}".rstrip(),
            f"- Quelle: {source}",
            f"- URL: {activity.url or '–'}",
        ]
        if activity.lv_number is not None:
            lines.append(f"- LV: {activity.lv_number}")
        return "\n".join(lines) + "\n\n---\n\n"

    def _save_text(self, activity: Activity, text: str, source: str) -> None:
        text = text.strip()
        if not text:
            return
        section = activity.section_number if activity.section_number is not None else 99
        path = self.root / "texte" / f"{section:02d}_{activity.index:03d}_{slug(activity.name)}.md"
        if path.exists() and activity.text_file:
            with path.open("a", encoding="utf-8") as handle:
                handle.write(f"\n\n## {source}\n\n{text}\n")
        else:
            path.write_text(self._front_matter(activity, source) + text + "\n", encoding="utf-8")
        activity.text_file = path.relative_to(self.root).as_posix()
        activity.characters += len(text)

    def _write_section_summaries(self, sections: dict[int | None, dict]) -> None:
        parts = []
        for number in sorted(sections, key=lambda n: (n is None, n)):
            section = sections[number]
            if section["summary"]:
                parts.append(f"## {number if number is not None else '–'} {section['name']}\n\n{section['summary']}")
        if parts:
            (self.root / "texte" / "00_abschnitte.md").write_text(
                "# Abschnittsbeschreibungen\n\n" + "\n\n".join(parts) + "\n", encoding="utf-8"
            )

    def _write_inventory(
        self, title: str, course_url: str, sections: dict, activities: list[Activity]
    ) -> None:
        created = datetime.now().isoformat(timespec="seconds")
        payload = {
            "course": {"title": title, "url": course_url, "exported_at": created},
            "sections": [sections[k] for k in sorted(sections, key=lambda n: (n is None, n))],
            "activities": [asdict(a) for a in activities],
        }
        (self.root / "inventar.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (self.root / "inventar.md").write_text(
            render_inventory_markdown(title, course_url, created, sections, activities), encoding="utf-8"
        )


def name_activities(activities: list[Activity], sections: dict, items: list[dict]) -> None:
    """Fill in section names and give untitled text fields a readable name.

    Section pages often list activities without the section heading around
    them; the heading is known from the section list instead. Text fields have
    no title of their own, so their first line serves as name.
    """
    for activity, item in zip(activities, items):
        section = sections.get(activity.section_number)
        if section and not activity.section_name:
            activity.section_name = section.get("name") or ""
        inline = (item.get("inline_text") or "").strip()
        if activity.name == "(ohne Titel)" and inline:
            first_line = inline.splitlines()[0].strip()
            activity.name = first_line[:80] + ("…" if len(first_line) > 80 else "")
        if activity.lv_number is None:
            activity.lv_number = lv_number(activity.name, activity.section_name, inline[:200])


def rebuild_report(root: Path) -> None:
    """Rewrite inventar.md (and names in inventar.json) from an existing export."""
    payload = json.loads((root / "inventar.json").read_text(encoding="utf-8"))
    sections = {s["number"]: s for s in payload["sections"]}
    activities = [Activity(**a) for a in payload["activities"]]
    items = []
    for activity in activities:
        inline = ""
        if activity.modtype == "label" and activity.text_file:
            inline = (root / activity.text_file).read_text(encoding="utf-8").split("\n---\n", 1)[-1]
        items.append({"inline_text": inline})
    name_activities(activities, sections, items)
    payload["activities"] = [asdict(a) for a in activities]
    (root / "inventar.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    course = payload["course"]
    (root / "inventar.md").write_text(
        render_inventory_markdown(course["title"], course["url"], course["exported_at"], sections, activities),
        encoding="utf-8",
    )


def render_inventory_markdown(
    title: str, course_url: str, created: str, sections: dict, activities: list[Activity]
) -> str:
    by_type: dict[str, int] = {}
    for activity in activities:
        by_type[activity.modtype] = by_type.get(activity.modtype, 0) + 1
    lv_map: dict[int, list[Activity]] = {}
    for activity in activities:
        if activity.lv_number is not None:
            lv_map.setdefault(activity.lv_number, []).append(activity)

    out = [
        f"# Kursinventar: {title}",
        "",
        f"Erfasst am {created} · {course_url}",
        "",
        f"**{len(activities)} Aktivitäten** in {len(sections)} Abschnitten · "
        + " · ".join(f"{t} {n}" for t, n in sorted(by_type.items(), key=lambda kv: -kv[1])),
        "",
        "## LV-Nummern",
        "",
    ]
    if lv_map:
        highest = max(lv_map)
        missing = [n for n in range(1, highest + 1) if n not in lv_map]
        out.append("Gefunden: " + ", ".join(str(n) for n in sorted(lv_map)))
        out.append("")
        out.append("Im Kurs nicht vorhanden (zwischen 1 und " + str(highest) + "): " + (", ".join(str(n) for n in missing) or "keine"))
        out.append("")
        out.append("| LV | Typ | Name | Erfasst |")
        out.append("|---|---|---|---|")
        for number in sorted(lv_map):
            for activity in lv_map[number]:
                out.append(f"| {number} | {activity.modtype} | {activity.name} | {describe(activity)} |")
    else:
        out.append("Keine Aktivität trägt eine LV-Nummer im Namen.")
    out += ["", "## Alle Aktivitäten nach Abschnitt", ""]
    current = object()
    for activity in activities:
        key = (activity.section_number, activity.section_name)
        if key != current:
            current = key
            label = f"{activity.section_number}" if activity.section_number is not None else "–"
            out += ["", f"### Abschnitt {label}: {activity.section_name or '(ohne Namen)'}", "", "| # | Typ | Name | Erfasst | Hinweis |", "|---|---|---|---|---|"]
        note = activity.note or activity.availability
        if activity.hidden:
            note = ("verborgen; " + note).strip("; ")
        out.append(f"| {activity.index} | {activity.modtype} | {activity.name} | {describe(activity)} | {note.replace('|', '/').replace(chr(10), ' ')} |")
    problems = [a for a in activities if a.status in {"fehler", "nicht geöffnet", "gesperrt", "ohne Text", "ohne Datei"}]
    if problems:
        out += ["", "## Nicht oder nicht vollständig erfasst", ""]
        for activity in problems:
            out.append(f"- **{activity.name}** ({activity.modtype}): {activity.status}. {activity.note}".rstrip())
    media = [(a, m) for a in activities for m in a.media]
    if media:
        out += ["", "## Medien (Audio, Video, eingebettete Inhalte)", ""]
        for activity, item in media:
            out.append(f"- {activity.name}: {item.get('tag')} {item.get('src')}")
    return "\n".join(out).rstrip() + "\n"


def describe(activity: Activity) -> str:
    parts = [activity.status]
    if activity.characters:
        parts.append(f"{activity.characters:,} Zeichen".replace(",", "."))
    if activity.files:
        parts.append(f"{len(activity.files)} Datei(en)")
    if activity.media:
        parts.append(f"{len(activity.media)} Medien")
    return ", ".join(parts)


def extract_file_text(path: Path) -> str:
    """Return the plain text of a PDF, PowerPoint or Word file; empty if unsupported."""
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="replace").strip()
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = [f"## Seite {i}\n\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages, start=1)]
            return "\n\n".join(pages).strip()
        if suffix == ".pptx":
            from pptx import Presentation

            slides = []
            for number, slide in enumerate(Presentation(str(path)).slides, start=1):
                texts = [shape.text_frame.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    texts.append("Notizen: " + slide.notes_slide.notes_text_frame.text.strip())
                slides.append(f"## Folie {number}\n\n" + "\n\n".join(texts))
            return "\n\n".join(slides).strip()
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            lines = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    lines.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
            return "\n\n".join(lines).strip()
    except ImportError:
        return ""
    except Exception as exc:
        return f"(Text konnte nicht gelesen werden: {exc})"
    return ""
